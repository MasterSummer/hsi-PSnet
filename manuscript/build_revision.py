"""Rebuild the evidence-based revision from the retained v6 Word source.

Usage: python manuscript/build_revision.py --source /path/to/v6.docx
Requires python-docx. Scientific summary CSVs and figure assets are versioned.
"""
import argparse
import csv
import difflib
import re
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent

REPLACEMENTS = {
1: 'Evaluating low cost hyperspectral and RGB imaging of inoculated Arabidopsis plants under controlled conditions',
11: 'Inoculated and mock plants compared within each dpi using plant-disjoint folds',
13: 'A benefit of hyperspectral information over RGB was not demonstrated',
16: 'Early detection of plant infection requires evaluation that separates treatment-associated signals from repeated observations and acquisition effects. We evaluated a Raspberry Pi hyperspectral imaging (HSI) unit paired with RGB imaging in a single controlled experiment with Arabidopsis thaliana inoculated with Albugo candida. The documented camera components cost £489.20, excluding printing, illumination and supporting equipment. The design comprised 96 plants (72 inoculated and 24 mock-inoculated) imaged at 2, 4 and 6 days post inoculation (dpi); the archived analysis used 569 paired observations. Date-matched binary tasks were evaluated with five plant-disjoint folds and three training seeds, with probabilities aggregated within plants. The dual-stream PSNet model achieved balanced accuracies of 56.16%, 83.56% and 95.37% at 2, 4 and 6 dpi, respectively, and corresponding AUROCs of 0.605, 0.939 and 0.990. RGB-only models performed within the range of multimodal models, and no architectural component showed a consistent benefit across tasks. Exploratory whole-cube spectral comparisons identified 302 of 313 computational bands at 2 dpi, but none at 4 or 6 dpi, after within-date false-discovery-rate correction. This reversed temporal pattern does not establish an infection-specific spectral signal. A separate single-seed, pooled-date experiment found that selecting ten bands on fitting plants retained high classification scores, but paired intervals did not establish an advantage over its full-band reference. Treatment was confounded with tray and acquisition order, and no independent cultivation run was available. The results demonstrate within-run group discrimination under controlled conditions and do not establish field performance or infection-specific diagnosis.',
27: 'Figure 1. Experimental design and analysis workflow. Arabidopsis thaliana plants were inoculated with Albugo candida or mock-inoculated and imaged at 2, 4 and 6 dpi. Paired HSI and RGB observations were partitioned by plant and evaluated in date-matched binary tasks. Microscopy was reported for a reserved leaf subset; it was not a quantitative per-plant reference standard for the classification labels.',
30: 'Spray guns, flasks and bottles were wiped with 100% ethanol before use. Leaves bearing visible A. candida pustules were suspended in chilled sterile water, and sporangia were dislodged by gentle agitation, filtered through nylon mesh and diluted on ice to 10⁵ spores ml⁻¹. Controls received sterile water only. A single run used 96 seedlings (72 inoculated, 24 mock-inoculated). Both leaf surfaces were misted to run-off. Trays were held at 4 °C for 24 h and then at 22 °C with 10 h light and 18 °C with 14 h dark until 2 dpi, after which standard growth conditions resumed. Symptoms were assessed visually; abaxial sporulation was observed around 6–8 dpi. Trypan blue staining and microscopy were reported for a reserved subset at 2, 4 and 6 dpi, following the previously cited workflow [24]. The available records do not establish stained-leaf denominators, matched mock results or a per-plant infection confirmation rate. Accordingly, the computational labels denote inoculation treatment, not independently confirmed infection in every plant.',
32: 'The hyperspectral camera was adapted from the open-source 3D-printed design of Salazar-Vazquez and Mendez-Vazquez [20]. The camera, controller and software were updated (Table 1). The camera mount, Part07_Camera, was redesigned as Part07_Camera_Module_3 to accommodate the different lens and autofocus geometry while retaining the PCB mounting-hole positions. The remaining enclosure and optical components were retained except for the objective-lens upgrade described below. Alternative case and lid variants were prototyped but were not used in the final system. The unit is an engineering adaptation of an existing design; the original design files are described in [20].',
37: 'The documented components total £489.20: the Edmund Optics VIS–NIR lens (£382.50), Raspberry Pi 4 Model B with 8 GB RAM (£62.50), Camera Module 3 NoIR (£21.70), and replacement-price estimates for the reused +10 close-up lens (£13.50) and grating (£9.00). Storage and power accessories were supplied with the Raspberry Pi. Printing costs and the price-record date were not retained, so this is a component subtotal rather than a verified total build cost. The laboratory illumination source, copy stand and cube-assembly workstation are excluded.',
42: 'The reconstruction workflow followed Salazar-Vazquez and Mendez-Vazquez [20], using fluorescent-lamp emission peaks, an incandescent-lamp spectrum and lens-covered dark frames for wavelength calibration and dark correction. The available records do not establish the timing of each calibration frame relative to individual trays. Archived model inputs contain 313 computational bands. Acquisition documentation reports 273 calibrated positions spanning 400.000–968.487 nm, but their correspondence to the historical PT tensor axis has not been independently verified. The subsequently recovered MAT files contain a nominal 313-position axis spanning 400.000–1052.090 nm; this alone does not validate the calibration range or the historical band mapping. Analyses therefore use computational band index. No white reference was used, and the data must not be interpreted as absolute reflectance. Spectral resolution, signal-to-noise ratio, linearity, repeatability, saturation behaviour and RGB–HSI registration accuracy were not independently characterised (Table 2).',
43: 'Archived HSI tensors have 313 channels and 132 × 135 spatial pixels. The retained training code reads these tensors directly; the upstream masking and intensity-transformation provenance is incomplete, so the spectral summaries are described as whole-cube rather than verified leaf-only measurements. Per-band standardisation statistics were estimated using only the fitting plants inside each outer fold, excluding both inner-validation and outer-test plants, and then applied unchanged. RGB images were resized to 224 × 224 pixels and normalised with the ImageNet channel means and standard deviations.',
47: 'The unit of analysis is the plant. Predicted positive-class probabilities were averaged over its available observations, including available dates in pooled tasks, and thresholded at 0.5; AUROC used the unthresholded means. Missing observations were not imputed. Five treatment-stratified, plant-disjoint outer folds were generated with split seed 157 and held fixed across models and training seeds within each task. Different tasks need not have identical fold assignments. Each outer-training partition was further split by plant into fitting and inner-validation sets, using a stratified 20% inner-validation fraction and random state equal to the training seed plus fold index. The checkpoint with the highest inner-validation plant-level balanced accuracy was retained. The outer-test fold was not used for checkpoint selection. The five held-out prediction sets were pooled to compute one plant-level metric per training seed (157, 257 or 357).',
55: 'Figure 3. Binary PSNet reimplementation evaluated in this revision. The RGB ResNet34 and HSI Transformer produce 64-dimensional embeddings. Concatenation and a 128 → 128 → 2 head yield mock-versus-inoculated logits. Tensor dimensions omit the batch axis. The model diagram and equations describe the released binary implementation rather than the earlier four-class prototype.',
59: 'The spectral encoder accepts X with shape C × H × W, where C = 313, H = 132 and W = 135. It constructs one global token from the spatial mean spectrum and 512 local tokens from learned feature channels (Fig. 4). All tokens have dimension 64.',
61: 'Figure 4. Implemented HSI encoder. A three-level Haar transform acts on the spatially averaged input spectrum. A separate 3 × 3 convolution reduces the cube to 64 learned feature channels, followed by 3D patch embedding. Five Transformer layers use four attention heads and CAF before layers 3–5. The feature-channel axis after reduction is not a physical wavelength axis. Supplementary Algorithm S1 specifies the global-token and CAF operations.',
62: 'Global token. The input is first averaged spatially. A three-level Haar transform, not a Daubechies-4 transform, is then applied to this mean spectrum. At each level, the final value is replicated if the length is odd. For paired values u and v, approximation and detail coefficients are (u + v)/√2 and (u − v)/√2, respectively:',
64: 'The detail vectors from levels 1, 2 and 3 are concatenated in that order, followed by the level-3 approximation. For 313 input bands, the concatenation has 316 elements; the implementation keeps the first 313 elements. A learned linear projection and GELU produce the global token:',
66: 'The resulting token has 64 elements and depends on the input cube. This differs from a freely learned CLS vector and from taking the mean of patch tokens. Spatial averaging precedes the Haar transform in the released implementation.',
67: 'Patch tokens. A Conv2D layer with 313 input channels, 64 output channels, kernel 3 × 3 and padding 1, followed by batch normalisation and GELU, produces a 64 × 132 × 135 feature tensor. A singleton channel axis is inserted for a Conv3D layer with 64 output channels, kernel and stride 8 × 16 × 16, and no padding. Thus:',
69: 'The 3D convolution spans blocks of eight learned feature channels and 16 × 16 spatial positions. Spatial remainders are discarded; the output grid is 8 × 8 × 8. These blocks do not represent groups of eight adjacent physical wavelengths after the learned channel reduction.',
71: 'Flattening this grid gives 512 patch tokens, each of dimension 64. The global token is prepended and positional embeddings are added to form the Transformer input:',
73: 'The positional parameter stores 1,025 × 64 elements, but only the first 513 positions are used in this path. The remaining 32,768 elements are registered but unused. Token dropout is 0.5.',
74: 'Encoder and cross-layer adaptive fusion. Five pre-normalisation Transformer layers use four attention heads, a feed-forward width of 256, GELU and dropout 0.5. CAF, adapted from SpectralFormer [25], concatenates the current representation with an earlier stored layer input and applies a learned 128 → 64 linear projection before layers 3, 4 and 5:',
76: 'Here h denotes the current layer representation and u denotes the stored layer input. The earlier inputs are those of layers 1, 2 and 3, respectively. The post-fusion input is stored before applying each Transformer layer. A final layer normalisation is applied, and the global-token position supplies the 64-dimensional HSI embedding.',
78: 'The RGB and HSI embeddings are concatenated into 128 elements. The classification head comprises Linear(128, 128), GELU, dropout 0.5 and Linear(128, 2). Training minimises ordinary, unweighted cross-entropy:',
80: 'There are two classes: mock (0) and inoculated (1). No class weighting or class-balanced resampling is implemented. Balanced accuracy and AUROC, rather than accuracy alone, are used to interpret the unequal treatment groups.',
82: 'Baselines use the same within-task outer folds and metrics. RGB-only ResNet18 and ResNet34 use ImageNet initialisation, global pooling, a 64-dimensional projection and a two-class head. The HSI Transformer uses the spectral branch defined above without RGB. The spectral 1D CNN averages each cube spatially and applies convolutional widths 32, 64 and 64 with kernels 7, 5 and 3, followed by global pooling and classification. The compact 3D CNN processes a singleton-channel cube with output widths 16, 32 and 64 and global pooling. The simple multimodal baseline combines ResNet18 and the spectral 1D CNN. The plain multimodal model uses a learnable CLS token and disables CAF. Historical PSNet variants disable CAF, substitute the mean patch token or a learnable CLS vector, or use direct 2D tokenisation. The latter changes capacity and token count as well as patching, so it is not an isolated test of 3D patching. Additional controlled substitutions are implemented in the public rerun code but have no completed real-data results in this manuscript.',
83: 'Table 4. Configurations and registered parameter counts at embedding dimension 64. Counts include 24,768 bypassed CAF parameters where CAF is disabled and 32,768 unused positional parameters in the full 3D path. Equal registered counts do not imply equal active capacity. Standalone baseline counts include their classification heads.',
85: 'Table 5 distinguishes settings recoverable from archived configurations from the documented rerun protocol. The archived configuration files verify embedding dimension 64, five Transformer layers, four attention heads, dropout 0.5 and pretrained RGB initialisation. The released training implementation uses AdamW, paired horizontal and vertical flips, inner-validation checkpoint selection and early stopping; it contains no learning-rate warm-up or scheduler, label smoothing or stochastic depth. Some historical run records do not retain every optimiser and training-budget argument, so the documented defaults cannot certify every historical invocation. New runs save the configuration, source snapshot, task hashes, fitted preprocessing, checkpoints and held-out predictions. No hyperparameter search is specified in the rerun protocol.',
86: 'Table 5. Audited architecture and documented rerun settings. Architecture settings are recovered from the archived configurations; optimiser and budget values specify the released rerun protocol where historical invocation records are incomplete.',
88: 'For exploratory spectral analysis, each archived cube was averaged over all spatial positions, and the available leaf means were averaged within each plant and date. No additional masking was assumed. At each computational band and separately within each dpi, inoculated and mock plants were compared with a two-sided Welch test. Benjamini–Hochberg correction was applied independently to the 313 tests at each date. Hedges’ g describes the inoculated-minus-mock standardised difference with small-sample correction. Figure 7 shows pointwise 95% t-based intervals for group means, not simultaneous confidence bands. The same plants recur across dates, so the three analyses are not independent. No red-edge wavelength, physiological attribution or gain-adjusted effect is claimed because the historical band mapping, masks and per-image gain records are not verified. Separately, an exploratory pooled-date experiment selected the 3, 10 or 30 bands with the largest absolute treatment-mean differences using fitting plants only inside each outer fold. Band selection and standardisation excluded inner-validation and outer-test plants. Each selection was retrained and compared with its own full-band control using seed 157 and five folds. Paired 95% percentile intervals used 20,000 treatment-stratified plant resamples of the fixed predictions, without retraining.',
93: 'Figure 5. Representative inoculated leaves at 2, 4 and 6 dpi with Albugo candida. White sporulation is visible on the abaxial surface at 6 dpi, whereas the imaged adaxial surface remains symptom-free. Scale bar, 2 mm. The panel documents representative inoculated material and does not provide a matched mock microscopy comparison or per-plant infection confirmation.',
101: 'Model rankings are descriptive. Between-seed variability and the conditional plant-bootstrap intervals in Table 11 indicate uncertainty, but neither constitutes a paired hypothesis test between models. Formal superiority or equivalence is not claimed. The HSI-only scores were highest at 2 dpi and close to chance later, including when abaxial sporulation was visible; this temporal pattern motivates the spectral and acquisition-confound analyses below.',
111: 'A separate, completed pooled-date band-selection experiment is summarised in Table 10. Its full-band control achieved 92.36% balanced accuracy; top-3, top-10 and top-30 inputs achieved 93.75%, 95.14% and 90.97%, respectively. Top-10 selection gave 96.88% accuracy but a lower AUROC than the full-band control (0.983 versus 0.991). The paired balanced-accuracy differences for top-3, top-10 and top-30 were +1.39, +2.78 and −1.39 percentage points, with 95% intervals of −6.25 to 9.03, −4.17 to 10.42 and −9.03 to 6.94. All include zero. These results show that reduced inputs can retain discrimination in this pooled task; they do not establish improved performance or a validated physiological band set. This full-band reference belongs to a separate experiment and must not be substituted for the earlier pooled result in Table 8.',
112: 'Table 10. Separate pooled-date band-selection experiment, seed 157 and five plant-disjoint folds. Selection uses fitting plants only. All dates include 6 dpi; these scores are not a measure of detection at 2 dpi. Each reduced-input model is compared with the full-band control from this experiment.',
117: 'At 2 dpi, 302 of 313 computational bands differed between groups at within-date BH q < 0.05; the minimum q was 0.000317 and the median absolute Hedges’ g was 0.786. At 4 and 6 dpi, no bands met this threshold; minimum q values were 0.0892 and 0.170, and median absolute effect sizes were 0.444 and 0.158, respectively (Table 12; Fig. 7). These are whole-cube intensity differences and can include background, illumination and acquisition effects. The absence of significant bands at later dates does not demonstrate absence of biological change. The particularly strong 2 dpi differences, together with the HSI-only classification pattern, warrant caution about an acquisition-specific signal.',
118: 'Table 12. Exploratory whole-cube spectral statistics. Welch tests and BH correction were performed separately for 313 computational bands within each dpi. Counts refer to available paired-data plants. Effect sizes are descriptive; bands are correlated and do not constitute independent biological replicates.',
119: 'Figure 7. Whole-cube spectral summaries at 2, 4 and 6 dpi. (a–c) Group mean intensity with pointwise 95% t-based confidence intervals over plant means. (d–f) Hedges’ g for inoculated minus mock plants; orange marks identify within-date BH q < 0.05. The x-axis denotes computational band index, not validated wavelength. Counts are 69 inoculated and 24 mock plants at 2 dpi, and 72 inoculated and 24 mock plants at 4 and 6 dpi.',
120: '3.9\tLimits of model attribution',
121: 'Archived band-group occlusion summaries reported no change in balanced accuracy or AUROC, but the corresponding masked probabilities were not retained and whole-modality interventions were not validated. We therefore do not use these summaries to infer that HSI is unused or to assign physiological meaning to particular bands. The band-selection experiment in Table 10 tests predictive sufficiency after retraining, which is distinct from attribution in an already trained model. No validated attribution result is presented.',
123: 'The full PSNet has 22,006,402 registered parameters (Table 4). Archived model-only inference on an NVIDIA A100-SXM4-80GB GPU with batch size 1 and synthetic input took 4.70 ms with peak GPU memory of 142.8 MiB. Acquisition and preprocessing are excluded. Comparable baseline timings, end-to-end latency and edge-device performance were not established, so these measurements cannot support a Raspberry Pi deployment claim.',
126: 'RGB-containing models distinguished the inoculated and mock groups at 4 dpi, before visible symptoms, and at 6 dpi, when sporulation was present on the lower surface but not on the imaged upper surface. PSNet was close to chance at 2 dpi. RGB-only models performed within the range of multimodal models, and the reported comparisons do not demonstrate an incremental benefit from HSI. The results establish within-run treatment-group discrimination, not confirmed infection-specific detection: tray, imaging order and unrecorded gain could also contribute.',
128: 'The HSI Transformer and compact 3D CNN achieved AUROCs of 0.850 and 0.916 at 2 dpi and then performed much less well at later dates. Whole-cube spectral differences were also strongest at 2 dpi. A transient biological response is one possible explanation, but acquisition-specific variation, background, automatic gain and overfitting remain alternatives. Mock plants were always imaged first, so treatment and acquisition order were inseparable. Neither the spectral summaries nor the classification scores identify the cause. The incomplete 2 dpi paired-data coverage and unverified upstream preprocessing further limit interpretation. We therefore do not treat the 2 dpi HSI-only scores as evidence of an infection-specific early signal.',
132: 'The historical ablations do not establish a consistent benefit of individual PSNet components. Rankings change across tasks, and the alternative 2D path changes token count and parameterisation together. The controlled raw-spectrum token, current-layer CAF and fixed depth-averaging substitutions in the released code would address some of these ambiguities, but they have not produced completed results for this revision. The separate top-k experiment addresses input dimensionality rather than these architectural mechanisms. Its paired intervals include zero, its band sets vary by fold, and it uses only one training seed and pooled dates.',
143: 'A low-cost adaptation of an open-source hyperspectral camera, paired with RGB imaging, was evaluated using date-matched comparisons and plant-disjoint folds. Inoculated and mock groups were distinguishable at 4 and 6 dpi in one controlled run, but PSNet was near chance at 2 dpi and a benefit of HSI beyond RGB was not demonstrated. Whole-cube spectral differences were strongest at 2 dpi and cannot be assigned an infection-specific cause. Reduced-band models retained high pooled-date scores without establishing superiority. Independent experimental runs with interleaved acquisition, verified plant identity, reflectance calibration and appropriate stress controls are required before these results can support presymptomatic diagnosis in greenhouse or field conditions.',
145: 'Proposed CRediT statement for author verification: Gabriel U. Crabb: Methodology (imaging and biological workflow), Investigation, Data curation, Formal analysis (exploratory modelling), Writing – original draft. Yiding Zhao: Software, Methodology (computational framework), Formal analysis (revised evaluation), Writing – review and editing. Xi Chen: Supervision, Methodology. Nicholas K. Priest: Supervision. Volkan Cevik: Conceptualization, Supervision. Final roles and approval remain to be confirmed by all authors.',
147: 'The competing-interest declaration remains subject to confirmation by all authors before submission.',
149: 'OpenAI Codex was used during preparation of this revision to assist with language editing, code organisation and checks against supplied analysis records. The authors must review the resulting text and code and take responsibility for the final submitted content. This disclosure is provided for author approval before submission.',
155: 'The computational code, revision manuscript and non-identifying summary tables are available at https://github.com/MasterSummer/hsi-PSnet. Raw paired images, cubes and model checkpoints are not included in the public revision package and may be requested from the corresponding authors. This manuscript reports the archived cohort and completed band-selection experiment; newly supplied raw files have not been incorporated into its classification results. The public recovery utility exports recovered MAT cubes and requires numerical agreement on overlapping observations before merging with historical inputs. Missing acquisition logs and incomplete upstream preprocessing provenance are described in the Methods.',
180: '[24] Furzer, O. J. et al. An improved assembly of the Albugo candida Ac2V genome reveals the expansion of the “CCG” class of effectors. Mol. Plant Microbe Interact. 35, 39–48 (2022). https://doi.org/10.1094/MPMI-04-21-0075-R',
195: 'Table S5. PSNet plant-level confusion counts per seed (positive class: inoculated). Counts for all listed tasks were recomputed from complete five-fold archived predictions, after averaging probabilities within plants and thresholding at 0.5. TP, inoculated plants correctly classified; FN, inoculated plants missed; TN, mock plants correctly classified; FP, mock plants classified as inoculated.'
}

FORMULAS = {
63: 's = mean_{HW}(X);  [d_{1}, d_{2}, d_{3}, a_{3}] = Haar_{3}(s)',
65: 'g = GELU(W_{g} · concat(d_{1}, d_{2}, d_{3}, a_{3})[0:313] + b_{g})',
68: 'N = floor(64/8) × floor(132/16) × floor(135/16) = 512',
70: 'T = flatten(Conv3D(reduce(X))) ∈ ℝ^{512×64}',
72: 'Z_{0} = Dropout(concat(g, T) + P[0:513]) ∈ ℝ^{513×64}',
75: 'u_{j} = W_{j} · concat(h_{j−1}, u_{j−2}) + b_{j},   j ∈ {3,4,5}',
79: 'L = −mean_{i} log softmax(logits_{i})[y_{i}]',
}


def set_text(p, text):
    # Retain paragraph formatting, replace wording and clear obsolete highlights.
    props = deepcopy(p._p.pPr) if p._p.pPr is not None else None
    p.clear()
    if props is not None:
        if p._p.pPr is not None:p._p.remove(p._p.pPr)
        p._p.insert(0, props)
    p.add_run(text)


def replace_table(table, rows):
    parent = table._tbl.getparent()
    doc = Document()
    new = doc.add_table(rows=0, cols=len(rows[0]))
    new.style = 'Table Grid'
    for i, values in enumerate(rows):
        cells = new.add_row().cells
        for c, value in zip(cells, values):
            c.text = str(value)
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(3)
                for r in p.runs:r.font.size=Pt(9);r.bold=(i==0)
        if i==0:
            repeat=OxmlElement('w:tblHeader');new.rows[0]._tr.get_or_add_trPr().append(repeat)
    parent.replace(table._tbl, new._tbl)


def extract(doc):
    lines=[]
    for item in doc.element.body:
        if item.tag==qn('w:p'):
            lines.append(''.join(item.xpath('.//w:t/text() | .//m:t/text()')))
        elif item.tag==qn('w:tbl'):
            for row in item.xpath('./w:tr'):
                lines.append(' | '.join(''.join(c.xpath('.//w:t/text() | .//m:t/text()')) for c in row.xpath('./w:tc')))
            lines.append('')
    return '\n'.join(lines)+'\n'


def build(source):
    doc=Document(source); before=extract(doc); pp=list(doc.paragraphs)
    assert len(pp)==195, 'Wrong source version: expected 195 top-level paragraphs'
    for n,text in REPLACEMENTS.items():set_text(pp[n-1],text)
    # Replace unsupported promises by bounded descriptions while preserving source details.
    for n in [33,39,91,136,139]:
        import re
        text=pp[n-1].text
        text=re.sub(r'\s*\[(?:ADD|CONFIRM|SUPERVISOR TO CONFIRM)[^\]]*\]', '',text)
        text=text.replace('The sub-£500 cost, of which the lens is about 78%, refers to the camera unit alone and excludes the laboratory-grade illumination source.', 'The £489.20 documented component subtotal excludes printing, illumination and supporting equipment.')
        text=text.replace('and its design files are open', 'and it derives from an open design')
        text=text.replace('(ix) Spectral statistics and model attribution are in progress.', '(ix) Whole-cube spectral statistics do not isolate leaf tissue or establish physiological causation. Validated model attribution, quantitative per-plant infection confirmation and an independent acquisition run are unavailable.')
        set_text(pp[n-1],text)
    # Remove the forced break that otherwise strands keywords on an empty page.
    pp[17].paragraph_format.page_break_before=False
    for break_node in pp[17]._p.xpath('.//w:br[@w:type="page"]'):
        break_node.getparent().remove(break_node)
    for number,(n,text) in enumerate(FORMULAS.items(),1):
        p=pp[n-1];p.clear()
        def mr(value):
            r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=value;r.append(t);return r
        math=OxmlElement('m:oMath');cursor=0
        for match in re.finditer(r'([A-Za-zℝ]+)([_^])\{([^}]+)\}',text):
            if match.start()>cursor:math.append(mr(text[cursor:match.start()]))
            node=OxmlElement('m:sSub' if match[2]=='_' else 'm:sSup')
            base=OxmlElement('m:e');base.append(mr(match[1]));node.append(base)
            sub=OxmlElement('m:sub' if match[2]=='_' else 'm:sup');sub.append(mr(match[3]));node.append(sub);math.append(node)
            cursor=match.end()
        math.append(mr(text[cursor:]));p._p.append(math)
        p.add_run(f'  ({number})')
    # Replace model diagrams with audited binary implementation diagrams.
    for n,filename in [(26,'Figure_1_workflow.png'),(54,'Figure_3_PSNet_overview.png'),(60,'Figure_4_HSI_encoder.png')]:
        p=pp[n-1];p.clear();p.add_run().add_picture(str(ROOT/'assets'/filename),width=Inches(6.15))
        p.paragraph_format.keep_with_next=True
    fig=pp[118].insert_paragraph_before()
    fig.add_run().add_picture(str(ROOT/'assets/Figure_7_spectral_statistics.png'),width=Inches(5.6))
    fig.alignment=1
    fig.paragraph_format.keep_with_next=True
    pp[118].paragraph_format.keep_together=True
    # Table numbering remains unchanged; Table 10 now contains completed results.
    tables=list(doc.tables)
    for row in tables[1].rows:
        for cell in row.cells:
            if 'sub-£500' in cell.text:cell.text=cell.text.replace('sub-£500 unit cost','£489.20 component subtotal')
            if 'Band-to-wavelength mapping' in cell.text:cell.text='Historical tensor-to-wavelength mapping unverified'
            if '[CONFIRM:' in cell.text:cell.text='Not independently logged'
    for name,count,detail in [('RGB ResNet18','11,209,474','RGB only'),('RGB ResNet34','21,317,634','RGB only'),('HSI Transformer','672,258','HSI branch only'),('Spectral 1D CNN','23,234','Spatial mean spectrum'),('Compact 3D CNN','81,474','Whole cube'),('Simple multimodal','11,240,834','ResNet18 + spectral 1D CNN')]:
        cells=tables[3].add_row().cells
        for cell,value in zip(cells,[name,'—','—',count,detail]):cell.text=value
    replace_table(tables[4],[['Item','Setting and evidence status'],
        ['RGB / HSI input','3 × 224 × 224 / 313 × 132 × 135'],
        ['Patch kernel and stride','8 × 16 × 16 after reduction to 64 learned channels'],
        ['Tokens / embedding','513 total / 64 dimensions'],
        ['Transformer depth / heads','5 / 4 (archived configuration)'],
        ['Classification head','128 → 128 → 2; GELU; dropout 0.5'],
        ['Optimiser','AdamW; learning rate 3 × 10⁻⁴; weight decay 10⁻⁴ (rerun protocol)'],
        ['Batch / maximum epochs','8 / 100 (rerun protocol)'],
        ['Checkpoint / early stopping','Inner-validation plant BA; patience 12 epochs; fraction 0.2'],
        ['Scheduler / loss','No scheduler or warm-up; unweighted cross-entropy'],
        ['Training seeds','157, 257, 357; outer split seed 157'],
        ['Augmentation / regularisation','Paired horizontal/vertical flips; dropout 0.5; no label smoothing or stochastic depth'],
        ['Hardware evidence','A100-SXM4-80GB for recorded latency; not a verified inventory of all training runs']])
    rows=list(csv.DictReader((ROOT/'evidence/band_comparison.csv').open()))
    pct=lambda value: str((Decimal(value)*100).quantize(Decimal('0.01'),rounding=ROUND_HALF_UP))
    replace_table(tables[9],[['Input','Bands','Accuracy (%)','BA (%)','AUROC','F1']]+[[r['band_set'],r['n_input_bands'],pct(r['accuracy']),pct(r['balanced_accuracy']),f"{float(r['auroc']):.3f}",f"{float(r['f1']):.3f}"] for r in rows])
    rows=list(csv.DictReader((ROOT/'evidence/spectral_summary.csv').open()))
    replace_table(tables[11],[['dpi','Inoculated / mock','Bands','BH q < 0.05','Minimum q','Median |g|']]+[[r['dpi'],f"{r['n_infected']} / {r['n_mock']}",r['n_bands'],r['significant_bands_within_dpi'],f"{float(r['minimum_q']):.6f}",f"{float(r['median_abs_hedges_g']):.3f}"] for r in rows])
    for row in tables[12].rows:
        for cell in row.cells:
            if 'sub-£500' in cell.text:cell.text=cell.text.replace('sub-£500 cost','£489.20 component subtotal')
    # Update Table S5 from verified archived predictions, not rounded rates.
    counts=list(csv.DictReader((ROOT/'evidence/verified_archived_metrics.csv').open()))
    names={'dpi_2':'2 dpi','dpi_4':'4 dpi','dpi_6':'6 dpi','presymptomatic_2_4':'Joint 2+4 dpi','all_dpi':'All dates'}
    replace_table(tables[17],[['Task','Seed','TP','FN','TN','FP']]+[[names[r['task']],r['seed'],r['TP'],r['FN'],r['TN'],r['FP']] for r in counts if r['model']=='psnet_full'])
    heading=doc.add_paragraph('Supplementary Algorithm S1',style=pp[185].style)
    heading.paragraph_format.keep_with_next=True
    doc.add_paragraph('Global token: average X over height and width; repeat three times: replicate an odd-length endpoint, compute Haar detail and approximation; concatenate [d1, d2, d3, a3], keep the first C entries, then apply Linear(C, 64) and GELU. For C = 313, coefficient lengths are 157, 79, 40 and 40 before truncation.')
    doc.add_paragraph('CAF: initialise history = []; for layer indices i = 0,...,4, first replace tokens with Linear(concat(tokens, history[i−2])) if i ≥ 2; append the resulting layer input to history; then apply Transformer layer i. Apply final LayerNorm and return token position 0. This records the post-fusion input, not the earlier layer output.')
    # Treatment wording does not assert that each inoculated plant was confirmed infected.
    for p in doc.paragraphs:
        for r in p.runs:
            revised=re.sub(r'\binfected\b','inoculated',r.text)
            revised=re.sub(r'\bInfected\b','Inoculated',revised)
            if revised != r.text:r.text=revised
            r.font.highlight_color=None
    # Keep existing document typography and table layout; remove stale yellow flags.
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(3)
                    p.paragraph_format.line_spacing=1
                    for r in p.runs:
                        r.font.highlight_color=None;r.font.size=Pt(9)
                        if r.text=='Infected':r.text='Inoculated'
    # Use a single page-number field in each footer (source contained duplicates).
    seen_footers=set()
    for section in doc.sections:
        footer=section.footer
        if id(footer._element) in seen_footers:continue
        seen_footers.add(id(footer._element))
        for child in list(footer._element):footer._element.remove(child)
        p=footer.add_paragraph();p.alignment=1
        suppress=OxmlElement('w:suppressLineNumbers');p._p.get_or_add_pPr().append(suppress)
        field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');p._p.append(field)
    doc.core_properties.comments='Evidence-based revision; author declarations require confirmation before submission.'
    doc.core_properties.last_modified_by=''
    output=ROOT/'PSNet_revision_20260928.docx';doc.save(output)
    after=extract(doc)
    (ROOT/'PSNet_revision_20260928.txt').write_text(after)
    (ROOT/'revision.diff').write_text(''.join(difflib.unified_diff(before.splitlines(True),after.splitlines(True),fromfile='v6_counts_verified',tofile='revision_20260928')))
    html=difflib.HtmlDiff(wrapcolumn=95).make_file(before.splitlines(),after.splitlines(),fromdesc='v6 counts verified',todesc='Revision 2026-09-28',context=True,numlines=2)
    (ROOT/'revision_diff.html').write_text('\n'.join(line.rstrip() for line in html.splitlines())+'\n')
    print(output)


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--source',type=Path,required=True)
    build(p.parse_args().source)
